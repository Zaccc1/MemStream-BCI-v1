import sys
import os
import datetime
import json
import numpy as np

import config


class TeeLogger:
    """Tee stdout to both terminal and log file."""

    def __init__(self, log_path):
        self.terminal = sys.stdout
        os.makedirs(os.path.dirname(log_path) or '.', exist_ok=True)
        self.log_file = open(log_path, 'w', encoding='utf-8')
        self.log_path = log_path

    def write(self, message):
        try:
            self.terminal.write(message)
        except UnicodeEncodeError:
            encoding = getattr(self.terminal, 'encoding', None) or 'utf-8'
            safe = message.encode(encoding, errors='replace').decode(encoding)
            self.terminal.write(safe)
        self.log_file.write(message)
        self.log_file.flush()

    def flush(self):
        self.terminal.flush()
        self.log_file.flush()

    def close(self):
        self.log_file.close()
        sys.stdout = self.terminal


def start_logging(output_dir="results"):
    """Start logging, return TeeLogger instance."""
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = os.path.join(output_dir, f"log_{timestamp}.txt")
    os.makedirs(output_dir, exist_ok=True)
    tee = TeeLogger(log_path)
    sys.stdout = tee
    print(f"[LOG] Logging to {log_path}")
    return tee


def stop_logging(tee):
    """Stop logging."""
    if tee is not None:
        print(f"[LOG] Log saved: {tee.log_path}")
        tee.close()


def _excel_safe(value):
    """Convert complex Python / NumPy objects into Excel-safe scalars/strings."""
    if value is None:
        return ''

    # Primitive Excel-safe scalars
    if isinstance(value, (str, int, float, bool, datetime.date, datetime.datetime)):
        return value

    # NumPy scalars
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating,)):
        return float(value)
    if isinstance(value, (np.bool_,)):
        return bool(value)

    # Arrays / sequences / mappings -> compact JSON/string
    if isinstance(value, np.ndarray):
        return json.dumps(value.tolist(), ensure_ascii=False)
    if isinstance(value, (list, tuple, set)):
        return json.dumps(list(value), ensure_ascii=False)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, default=str)

    # Fallback
    return str(value)


def collect_config_params():
    """Collect all uppercase parameters from config.py as a dict."""
    params = {}
    for key in sorted(dir(config)):
        if key.startswith('_'):
            continue
        if not key[0].isupper():
            continue
        val = getattr(config, key)
        params[key] = _excel_safe(val)
    return params


def collect_results(clust_results=None, sa_events=None, events=None,
                    recording_info=None, elapsed_s=None, comparison_summary=None):
    """Extract key metrics from run outputs."""
    results = {}
    results['timestamp'] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    if elapsed_s is not None:
        results['elapsed_s'] = round(elapsed_s, 1)

    if recording_info:
        results['source'] = recording_info.get('source', '?')
        results['total_duration_s'] = recording_info.get('total_duration_s', 0)
        results['calib_duration_s'] = recording_info.get('calibration_duration_s', 0)
        results['test_duration_s'] = recording_info.get('test_duration_s', 0)
        results['n_test_chunks'] = recording_info.get('n_test_chunks', 0)
        results['n_channels'] = recording_info.get('n_channels', 0)

    if events is not None:
        results['tm_n_spikes'] = int(len(events.get('times', [])))

    if sa_events is not None:
        n_sa = len(sa_events.get('times', []))
        results['sa_n_spikes'] = int(n_sa)
        if n_sa > 0:
            amps = sa_events.get('amplitudes', np.array([]))
            if len(amps) > 0:
                results['sa_amp_mean'] = round(float(np.mean(np.abs(amps))), 2)
                results['sa_amp_max'] = round(float(np.max(np.abs(amps))), 2)
            results['sa_active_ch'] = int(len(np.unique(sa_events.get('channels', []))))

    if clust_results is not None:
        results['clust_n_spikes'] = int(len(clust_results.get('times', [])))
        results['clust_n_active'] = int(clust_results.get('n_active', 0))
        results['clust_n_total'] = int(clust_results.get('n_clusters_total', 0))
        results['clust_mode'] = clust_results.get('mode', '?')
        results['clust_n_regions'] = int(clust_results.get('n_regions', 0))
        results['n_merges'] = int(clust_results.get('n_merges', 0))
        results['pre_merge_clusters'] = int(clust_results.get('pre_merge_n_clusters', 0))

        if 'is_good' in clust_results:
            results['n_good_units'] = int(clust_results['is_good'].sum())

        pca_var = clust_results.get('pca_variance_ratio', None)
        if pca_var is not None and len(pca_var) > 0:
            results['pca_cumvar_pct'] = round(float(np.sum(pca_var)) * 100, 1)

    if comparison_summary:
        for key, val in comparison_summary.items():
            if isinstance(val, (np.floating, float)):
                results[key] = round(float(val), 6)
            elif isinstance(val, (np.integer, int)):
                results[key] = int(val)
            else:
                results[key] = _excel_safe(val)

    return results


def append_to_xlsx(params, results, xlsx_path="results_log.xlsx"):
    """Append parameters and results as one row in xlsx. Creates file if needed."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    row_data = {}
    row_data.update({k: _excel_safe(v) for k, v in results.items()})
    row_data.update({k: _excel_safe(v) for k, v in params.items()})

    if os.path.exists(xlsx_path):
        wb = load_workbook(xlsx_path)
        ws = wb.active
        existing_headers = [cell.value for cell in ws[1]]
        for key in row_data:
            if key not in existing_headers:
                existing_headers.append(key)
                col_idx = len(existing_headers)
                ws.cell(row=1, column=col_idx, value=key)
                ws.cell(row=1, column=col_idx).font = Font(bold=True)
                ws.cell(row=1, column=col_idx).fill = PatternFill('solid', fgColor='D9E1F2')
        headers = existing_headers
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Run Log"
        headers = list(row_data.keys())
        for col_idx, key in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=key)
            cell.font = Font(bold=True, size=10)
            cell.fill = PatternFill('solid', fgColor='D9E1F2')
            cell.alignment = Alignment(horizontal='center')
        ws.freeze_panes = 'A2'

    new_row = ws.max_row + 1
    for col_idx, key in enumerate(headers, 1):
        val = row_data.get(key, '')
        ws.cell(row=new_row, column=col_idx, value=val)

    for col_idx, key in enumerate(headers, 1):
        max_len = max(len(str(key)), 8)
        for row in range(max(2, new_row - 3), new_row + 1):
            cell_val = ws.cell(row=row, column=col_idx).value
            if cell_val is not None:
                max_len = max(max_len, min(len(str(cell_val)), 30))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max_len + 2

    wb.save(xlsx_path)
    n_runs = ws.max_row - 1
    print(f"[LOG] Results appended to {xlsx_path} (run #{n_runs})")
    return xlsx_path
